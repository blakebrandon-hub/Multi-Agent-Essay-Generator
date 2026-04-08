import json
import os
import re
from flask import Flask, render_template, request, jsonify, send_file
from openai import OpenAI
import threading
from datetime import datetime
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ---------------------------------------------------
# Configuration
# ---------------------------------------------------

MODEL = "gpt-4o-mini"
MAX_ITERATIONS = 3
OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "outputs")
os.makedirs(OUTPUT_DIR, exist_ok=True)

app = Flask(__name__)
client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

# Store job results in memory (in production, use Redis or database)
job_results = {}

# ---------------------------------------------------
# LLM Wrapper
# ---------------------------------------------------

def call_llm(system_prompt, user_prompt, temperature=0.3):
    response = client.chat.completions.create(
        model=MODEL,
        temperature=temperature,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ],
    )
    return response.choices[0].message.content.strip()


# ---------------------------------------------------
# Agent Base Class
# ---------------------------------------------------

class Agent:
    def __init__(self, name, system_prompt):
        self.name = name
        self.system_prompt = system_prompt

    def run(self, prompt, temperature=0.3):
        return call_llm(self.system_prompt, prompt, temperature)


# ---------------------------------------------------
# Research Planner Agent
# ---------------------------------------------------

planner = Agent(
    "research_planner",
"""
You are a research planning agent for academic technical essays.

Given the essay topic and requirements, break down the research into specific, actionable research tasks.

Consider:
- Key concepts that need to be researched
- Technical background information needed
- Current state of the field
- Key arguments or perspectives
- Supporting evidence and data
- Counterarguments if applicable

Return JSON ONLY in this format:

{
  "research_tasks": [
    "Research task 1 - specific area to investigate",
    "Research task 2 - specific area to investigate",
    "Research task 3 - specific area to investigate"
  ],
  "essay_structure": [
    "Introduction",
    "Section 1 title",
    "Section 2 title",
    "Conclusion"
  ]
}
"""
)


# ---------------------------------------------------
# Deep Research Agent
# ---------------------------------------------------

researcher = Agent(
    "deep_researcher",
"""
You are an expert research agent specializing in technical and academic topics.

For each research task provided, conduct thorough research and provide:
- Key concepts and definitions
- Technical details and explanations
- Current developments and state of the field
- Important facts, figures, and data
- Expert perspectives and scholarly viewpoints
- Relevant examples and case studies

Organize your research clearly with headings for each research task.

Provide detailed, accurate, technical information suitable for a college-level essay.

Focus on depth and accuracy. Cite specific examples and technical details.
"""
)


# ---------------------------------------------------
# Outline Creator Agent
# ---------------------------------------------------

outliner = Agent(
    "outliner",
"""
You are an academic writing expert that creates detailed essay outlines.

Based on the topic, research findings, and suggested structure, create a comprehensive outline.

For each section:
- Provide a clear thesis or main point
- List 3-5 key points to cover
- Note which research findings support each point
- Suggest transitions between sections

Return a detailed outline in markdown format with clear hierarchical structure.

Use this format:
# Essay Title

## I. Introduction
- Hook/opening
- Background context
- Thesis statement

## II. [Section Title]
- Main point 1
  - Supporting detail from research
- Main point 2
  - Supporting detail from research

... and so on
"""
)


# ---------------------------------------------------
# Essay Writer Agent
# ---------------------------------------------------

writer = Agent(
    "essay_writer",
"""
You are an expert academic writer specializing in technical essays at the college level.

Write a complete, polished technical essay based on the outline and research provided.

Requirements:
- Clear, academic writing style
- Well-structured paragraphs with topic sentences
- Smooth transitions between ideas
- Technical accuracy
- Proper integration of research findings
- Strong thesis and conclusion
- College-level vocabulary and sophistication
- 1500-2500 words typical length

Write the COMPLETE essay, not an outline or summary.

Format with clear section headings using markdown:
# Title
## Section Headings
### Subsections if needed

Write full paragraphs with detailed explanations.
"""
)


# ---------------------------------------------------
# Editor Agent (Critic)
# ---------------------------------------------------

editor = Agent(
    "editor",
"""
You are an academic editor and writing instructor.

Evaluate the essay for:
1. **Thesis clarity** - Is there a clear, arguable thesis?
2. **Organization** - Does the essay follow a logical structure?
3. **Evidence** - Are claims well-supported with research?
4. **Technical accuracy** - Is technical information correct and precise?
5. **Writing quality** - Is the writing clear, engaging, and academic?
6. **Depth** - Does the essay demonstrate deep understanding?
7. **Completeness** - Are all required sections present and fully developed?

Return JSON ONLY:

{
  "pass": true or false,
  "score": {
    "thesis": 0-10,
    "organization": 0-10,
    "evidence": 0-10,
    "technical_accuracy": 0-10,
    "writing_quality": 0-10,
    "depth": 0-10
  },
  "strengths": ["strength 1", "strength 2"],
  "improvements_needed": ["improvement 1", "improvement 2"],
  "feedback": "overall assessment and specific suggestions"
}

Pass = true only if the essay is well-developed, well-researched, and meets college standards.
Pass = false if major revisions are needed.
"""
)


# ---------------------------------------------------
# Document Creator
# ---------------------------------------------------

def create_word_document(essay_content, metadata, job_id):
    """
    Create a properly formatted Word document from the essay markdown.
    """
    doc = Document()
    
    # Set document margins (1 inch all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    # Parse markdown and add to document
    lines = essay_content.split('\n')
    
    for line in lines:
        line = line.strip()
        
        if not line:
            continue
        
        # Title (# heading)
        if line.startswith('# '):
            title = line[2:].strip()
            p = doc.add_paragraph(title)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.runs[0]
            run.font.size = Pt(16)
            run.font.bold = True
            doc.add_paragraph()  # Space after title
            
        # Section heading (## heading)
        elif line.startswith('## '):
            heading = line[3:].strip()
            p = doc.add_paragraph(heading)
            run = p.runs[0]
            run.font.size = Pt(14)
            run.font.bold = True
            doc.add_paragraph()  # Space after heading
            
        # Subsection (### heading)
        elif line.startswith('### '):
            subheading = line[4:].strip()
            p = doc.add_paragraph(subheading)
            run = p.runs[0]
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.italic = True
            
        # Regular paragraph
        else:
            # Remove markdown bold/italic markers for cleaner text
            clean_line = re.sub(r'\*\*([^*]+)\*\*', r'\1', line)
            clean_line = re.sub(r'\*([^*]+)\*', r'\1', clean_line)
            
            p = doc.add_paragraph(clean_line)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.runs[0]
            run.font.size = Pt(12)
            run.font.name = 'Times New Roman'
    
    # Save document
    output_folder = os.path.join(OUTPUT_DIR, job_id)
    os.makedirs(output_folder, exist_ok=True)
    
    # Create safe filename
    topic_clean = re.sub(r'[^\w\s-]', '', metadata.get('topic', 'essay'))[:50]
    topic_clean = re.sub(r'[-\s]+', '_', topic_clean)
    
    filepath = os.path.join(output_folder, f"{topic_clean}.docx")
    doc.save(filepath)
    
    return filepath


# ---------------------------------------------------
# Orchestrator
# ---------------------------------------------------

def run_essay_agent(topic, requirements, job_id):
    """
    Run the essay research and writing pipeline.
    """
    
    state = {
        "topic": topic,
        "requirements": requirements,
        "research_tasks": [],
        "structure": [],
        "research": "",
        "outline": "",
        "essay": "",
        "feedback": "",
        "file": None,
        "status": "running",
        "current_iteration": 0,
        "current_phase": "",
        "logs": [],
        "scores": {}
    }
    
    job_results[job_id] = state

    try:
        # Phase 1: Planning
        state["current_phase"] = "Planning Research"
        state["logs"].append("🎯 Planning research approach...")
        
        planning_prompt = f"""
Topic: {topic}

Requirements:
{requirements}

Create a research plan and suggested essay structure.
"""
        
        plan_response = planner.run(planning_prompt)
        
        try:
            plan_data = json.loads(plan_response)
            state["research_tasks"] = plan_data.get("research_tasks", [])
            state["structure"] = plan_data.get("essay_structure", [])
            state["logs"].append(f"✓ Identified {len(state['research_tasks'])} research areas")
        except Exception as e:
            state["logs"].append(f"✗ Planning failed: {str(e)}")
            state["status"] = "error"
            return
        
        # Phase 2: Deep Research
        state["current_phase"] = "Conducting Research"
        state["logs"].append("📚 Conducting deep research...")
        
        research_prompt = f"""
Topic: {topic}

Research these specific areas:
{chr(10).join(f"- {task}" for task in state["research_tasks"])}

Requirements to keep in mind:
{requirements}

Provide comprehensive research for each area.
"""
        
        state["research"] = researcher.run(research_prompt, temperature=0.4)
        state["logs"].append("✓ Research completed")
        
        # Phase 3: Outline Creation
        state["current_phase"] = "Creating Outline"
        state["logs"].append("📝 Creating detailed outline...")
        
        outline_prompt = f"""
Topic: {topic}

Requirements:
{requirements}

Suggested Structure:
{chr(10).join(state["structure"])}

Research Findings:
{state["research"]}

Create a detailed outline integrating this research.
"""
        
        state["outline"] = outliner.run(outline_prompt)
        state["logs"].append("✓ Outline created")
        
        # Phase 4: Essay Writing (with iterations)
        for iteration in range(MAX_ITERATIONS):
            state["current_iteration"] = iteration + 1
            state["current_phase"] = f"Writing Essay (Draft {iteration + 1})"
            state["logs"].append(f"✍️ Writing draft {iteration + 1}...")
            
            writing_prompt = f"""
Topic: {topic}

Requirements:
{requirements}

Outline:
{state["outline"]}

Research:
{state["research"]}

{"Previous feedback to address:" + state["feedback"] if state["feedback"] else ""}

Write the complete essay now.
"""
            
            state["essay"] = writer.run(writing_prompt, temperature=0.5)
            state["logs"].append(f"✓ Draft {iteration + 1} completed")
            
            # Phase 5: Editorial Review
            state["current_phase"] = "Editorial Review"
            state["logs"].append("🔍 Editorial review in progress...")
            
            review_prompt = f"""
Topic: {topic}

Requirements:
{requirements}

Essay:
{state["essay"]}

Evaluate this essay thoroughly.
"""
            
            review_response = editor.run(review_prompt)
            
            try:
                review_data = json.loads(review_response)
                state["scores"] = review_data.get("score", {})
                
                if review_data.get("pass", False):
                    state["logs"].append("✅ Essay approved by editor!")
                    state["feedback"] = review_data.get("feedback", "")
                    break
                else:
                    improvements = review_data.get("improvements_needed", [])
                    state["logs"].append(f"⚠️ Revisions needed: {len(improvements)} issues")
                    state["feedback"] = review_data.get("feedback", "")
                    
                    if iteration < MAX_ITERATIONS - 1:
                        state["logs"].append("🔄 Preparing revision...")
            except Exception as e:
                state["logs"].append(f"✗ Review parsing error: {str(e)}")
                break
        
        # Phase 6: Create Word Document
        state["current_phase"] = "Creating Document"
        state["logs"].append("📄 Creating Word document...")
        
        metadata = {
            "topic": topic,
            "requirements": requirements
        }
        
        state["file"] = create_word_document(state["essay"], metadata, job_id)
        state["logs"].append("✓ Word document created")
        
        state["status"] = "completed"
        state["current_phase"] = "Complete"
        state["logs"].append("🎉 Essay complete!")
        
    except Exception as e:
        state["status"] = "error"
        state["logs"].append(f"❌ Error: {str(e)}")


# ---------------------------------------------------
# Flask Routes
# ---------------------------------------------------

@app.route('/')
def index():
    return render_template('index.html')


@app.route('/api/generate', methods=['POST'])
def generate():
    """Start a new essay generation job"""
    data = request.json
    topic = data.get('topic', '')
    requirements = data.get('requirements', '')
    
    if not topic:
        return jsonify({"error": "No topic provided"}), 400
    
    # Generate unique job ID
    job_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    # Start agent in background thread
    thread = threading.Thread(target=run_essay_agent, args=(topic, requirements, job_id))
    thread.start()
    
    return jsonify({"job_id": job_id})


@app.route('/api/status/<job_id>')
def get_status(job_id):
    """Get the current status of a job"""
    if job_id not in job_results:
        return jsonify({"error": "Job not found"}), 404
    
    state = job_results[job_id]
    
    return jsonify({
        "status": state["status"],
        "current_iteration": state["current_iteration"],
        "current_phase": state["current_phase"],
        "logs": state["logs"],
        "research_tasks": state["research_tasks"],
        "structure": state["structure"],
        "has_file": state["file"] is not None,
        "essay_preview": state["essay"][:500] if state["essay"] else "",
        "scores": state["scores"],
        "feedback": state["feedback"]
    })


@app.route('/api/download/<job_id>')
def download_file(job_id):
    """Download the generated essay"""
    if job_id not in job_results:
        return jsonify({"error": "Job not found"}), 404
    
    state = job_results[job_id]
    
    if not state["file"]:
        return jsonify({"error": "No file generated"}), 404
    
    return send_file(
        state["file"],
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        as_attachment=True,
        download_name=os.path.basename(state["file"])
    )


@app.route('/api/essay/<job_id>')
def get_essay(job_id):
    """Get the full essay text"""
    if job_id not in job_results:
        return jsonify({"error": "Job not found"}), 404
    
    state = job_results[job_id]
    
    return jsonify({
        "essay": state["essay"],
        "outline": state["outline"],
        "research": state["research"]
    })


if __name__ == "__main__":
    print("\n" + "="*60)
    print("Technical Essay Research Agent")
    print("="*60)
    print(f"Output directory: {OUTPUT_DIR}")
    print("Starting server on http://localhost:5000")
    print("="*60 + "\n")
    
    app.run(debug=True, host='0.0.0.0', port=5000)
